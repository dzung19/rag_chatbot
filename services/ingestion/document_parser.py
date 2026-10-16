"""Document parsers for supported file types.

Uses lazy-loading generator pattern for memory efficiency.
Supports: PDF, DOCX, PPTX, XLSX, TXT, Markdown.
"""

from __future__ import annotations

import logging
from typing import Generator

logger = logging.getLogger(__name__)


def parse_document(file_path: str, file_type: str) -> Generator[str, None, None]:
    """Parse a document and yield text content page-by-page.

    Args:
        file_path: Path to the document file.
        file_type: File type string (pdf, docx, pptx, xlsx, txt, md).

    Yields:
        Text content for each logical page/section of the document.
    """
    parser_map = {
        "pdf": _parse_pdf,
        "docx": _parse_docx,
        "pptx": _parse_pptx,
        "xlsx": _parse_xlsx,
        "txt": _parse_text,
        "md": _parse_text,
    }

    parser = parser_map.get(file_type)
    if parser is None:
        logger.error("No parser available for file type: %s", file_type)
        return

    yield from parser(file_path)


def _parse_pdf(file_path: str) -> Generator[str, None, None]:
    """Parse PDF using layout-aware & box-grouping PDF parser. Yields structured text per page."""
    try:
        from pdf_parser import parse_pdf_document

        yield from parse_pdf_document(file_path)
    except Exception as e:
        logger.error("Failed to parse PDF %s: %s", file_path, str(e))
        raise


def _parse_docx(file_path: str) -> Generator[str, None, None]:
    """Parse DOCX using python-docx. Yields paragraphs and tables."""
    from docx import Document

    try:
        doc = Document(file_path)

        # Extract paragraphs
        paragraphs = []
        for para in doc.paragraphs:
            if para.text and para.text.strip():
                paragraphs.append(para.text.strip())

        if paragraphs:
            yield "\n".join(paragraphs)

        # Extract tables
        for table_idx, table in enumerate(doc.tables):
            table_text_lines = []
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells]
                table_text_lines.append(" | ".join(row_data))
            if table_text_lines:
                yield f"[Table {table_idx + 1}]\n" + "\n".join(table_text_lines)

    except Exception as e:
        logger.error("Failed to parse DOCX %s: %s", file_path, str(e))
        raise


def _parse_pptx(file_path: str) -> Generator[str, None, None]:
    """Parse PPTX using python-pptx. Yields text per slide."""
    from pptx import Presentation

    try:
        prs = Presentation(file_path)

        for slide_idx, slide in enumerate(prs.slides, 1):
            slide_texts = []

            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            slide_texts.append(text)

                # Extract text from tables in slides
                if shape.has_table:
                    for row in shape.table.rows:
                        row_data = [cell.text.strip() for cell in row.cells]
                        row_text = " | ".join(row_data)
                        if row_text.strip():
                            slide_texts.append(row_text)

            # Include slide notes if present
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    slide_texts.append(f"[Notes] {notes}")

            if slide_texts:
                yield f"[Slide {slide_idx}]\n" + "\n".join(slide_texts)

    except Exception as e:
        logger.error("Failed to parse PPTX %s: %s", file_path, str(e))
        raise


def _parse_xlsx(file_path: str) -> Generator[str, None, None]:
    """Parse XLSX using dedicated ExcelParser. Yields clean markdown tables per sheet."""
    try:
        from excel_parser import parse_and_chunk_excel
        text_pages, _ = parse_and_chunk_excel(file_path)
        for page in text_pages:
            yield page
    except Exception as e:
        logger.error("Failed to parse XLSX %s: %s", file_path, str(e))
        raise


def _parse_text(file_path: str) -> Generator[str, None, None]:
    """Parse plain text or Markdown files."""
    encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252"]

    for encoding in encodings:
        try:
            with open(file_path, "r", encoding=encoding) as f:
                content = f.read()
            if content and content.strip():
                yield content.strip()
            return
        except (UnicodeDecodeError, UnicodeError):
            continue

    logger.error("Failed to decode text file %s with any encoding", file_path)
    raise ValueError(f"Unable to decode file: {file_path}")
