"""Diagnostic tool to inspect Excel files and estimate chunk counts.

Usage (on host or inside container):
    python scratch/analyze_excel.py path/to/document.xlsx
"""

import sys
from pathlib import Path

# Add shared and ingestion service to path
base_dir = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(base_dir))
sys.path.insert(0, str(base_dir / "services" / "ingestion"))

try:
    from openpyxl import load_workbook
    from excel_parser import ExcelParser
except ImportError as e:
    print(f"Error: Required dependency missing ({e}). Run inside rag-ingestion container or install openpyxl.")
    sys.exit(1)


def analyze_file(file_path: str):
    p = Path(file_path)
    if not p.exists():
        print(f"File not found: {file_path}")
        sys.exit(1)

    print("=" * 70)
    print(f"  EXCEL ANALYSIS REPORT: {p.name}")
    print(f"  Size: {p.stat().st_size / (1024 * 1024):.2f} MB")
    print("=" * 70)

    wb = load_workbook(file_path, read_only=True, data_only=True)
    try:
        sheet_count = len(wb.sheetnames)
        hidden_count = sum(
            1 for name in wb.sheetnames if getattr(wb[name], "sheet_state", "visible") != "visible"
        )
        print(f"\nTổng số sheet: {sheet_count} (Sheet ẩn: {hidden_count})")
        print("-" * 70)

        for idx, name in enumerate(wb.sheetnames, start=1):
            ws = wb[name]
            state = getattr(ws, "sheet_state", "visible")
            status_str = f"[{state.upper()}]" if state != "visible" else "[VISIBLE]"

            row_count = 0
            max_col_detected = 0
            non_empty_cells = 0
            for r in ws.iter_rows(values_only=True):
                row_count += 1
                row_non_empty = [c for c in r if c is not None and str(c).strip() != ""]
                non_empty_cells += len(row_non_empty)
                if row_non_empty:
                    last_idx = max(i for i, c in enumerate(r) if c is not None and str(c).strip() != "") + 1
                    if last_idx > max_col_detected:
                        max_col_detected = last_idx

            total_cells = row_count * max_col_detected if max_col_detected else 0
            fill_pct = (non_empty_cells / total_cells * 100) if total_cells else 0.0

            print(f"Sheet {idx}: '{name}' {status_str}")
            print(f"  * Số dòng quét: {row_count:,}")
            print(f"  * Số cột thực tế có dữ liệu: {max_col_detected}")
            print(f"  * Tỉ lệ ô có dữ liệu: {fill_pct:.1f}% ({non_empty_cells:,} / {total_cells:,} ô)")
    finally:
        wb.close()

    print("\n" + "=" * 70)
    print("  SIMULATION: PARSING & CHUNKING (ExcelParser v2)")
    print("=" * 70)

    parser = ExcelParser(
        chunk_chars=4000,
        full_index_rows=2000,
        max_chunks_per_file=500,
        table_gap_rows=2,
        sparse_threshold=0.4,
    )

    text_pages, chunks = parser.parse_and_chunk(file_path, source_filename=p.name)

    print(f"\n>> Kết quả sau khi tối ưu:")
    print(f"   * Tổng số text pages (SQLite full text): {len(text_pages)}")
    print(f"   * Tổng số vector chunks (Qdrant index): {len(chunks)}")

    summary_chunks = [c for c in chunks if c.get("metadata", {}).get("is_summary")]
    data_chunks = [c for c in chunks if not c.get("metadata", {}).get("is_summary")]
    print(f"     - Chunk tóm tắt / thống kê (high-level): {len(summary_chunks)}")
    print(f"     - Chunk dữ liệu chi tiết (row-window): {len(data_chunks)}")

    if chunks:
        print("\n--- Mẫu Chunk đầu tiên ---")
        first = chunks[0]
        print(f"Heading: {first['heading']}")
        print(f"Context: {first['context_header']}")
        print(f"Preview text:\n{first['text'][:300]}...")

    print("\n" + "=" * 70)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python scratch/analyze_excel.py <path_to_excel_file>")
        sys.exit(1)
    analyze_file(sys.argv[1])
